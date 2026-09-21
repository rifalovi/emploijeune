"""Lecture / écriture de fichiers SPSS (.sav) et construction d'un SurveyDataset.

Porté depuis `load` et `export_sav` de l'application desktop V4.8. pyreadstat
est importé paresseusement : les autres formats (CSV, Excel Kobo/CSPro) et tout
le reste du moteur fonctionnent sans lui.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

import pandas as pd

from .dataset import SurveyDataset
from .specs import infer_variable_specs


def _pyreadstat():
    try:
        import pyreadstat  # import paresseux
    except ImportError as exc:  # pragma: no cover - dépend de l'environnement
        raise RuntimeError(
            "Le module pyreadstat est requis pour les fichiers .sav : pip install pyreadstat"
        ) from exc
    return pyreadstat


def read_sav(path: str, name: str = "Données SPSS") -> SurveyDataset:
    """Lit un fichier .sav (codes bruts + métadonnées SPSS) en SurveyDataset."""
    df, meta = _pyreadstat().read_sav(path, apply_value_formats=False)
    return SurveyDataset(
        frame=df,
        variable_labels=meta.column_names_to_labels,
        value_labels=meta.variable_value_labels,
        variable_measure=getattr(meta, "variable_measure", {}) or {},
        original_variable_types=getattr(meta, "original_variable_types", {}) or {},
        name=name,
    )


# Extensions tabulaires prises en charge (hors .sav), pour l'UI et les messages.
TABULAR_EXTENSIONS = (
    ".xlsx",
    ".xls",
    ".ods",
    ".csv",
    ".tsv",
    ".tab",
    ".json",
    ".docx",
)


def _read_docx_table(path: str) -> pd.DataFrame:
    """Extrait le premier tableau d'un document Word (.docx) en DataFrame.

    La première ligne du tableau sert d'en-tête. Utile pour les résultats
    d'enquête présentés sous forme de tableau dans un document Word.
    """
    try:
        import docx  # python-docx, import paresseux
    except ImportError as exc:  # pragma: no cover - dépend de l'environnement
        raise RuntimeError(
            "Le module python-docx est requis pour les fichiers .docx : pip install python-docx"
        ) from exc
    document = docx.Document(path)
    if not document.tables:
        raise ValueError(
            "Ce document Word ne contient aucun tableau exploitable. "
            "Fournissez un tableau (ou exportez les données en Excel/CSV)."
        )
    table = document.tables[0]
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if len(rows) < 2:
        raise ValueError("Le tableau du document Word est vide (pas de données sous l'en-tête).")
    header = rows[0]
    return pd.DataFrame(rows[1:], columns=header)


def _excel_engine(ext: str) -> str:
    if ext == ".ods":
        return "odf"
    return "openpyxl" if ext == ".xlsx" else "xlrd"


def list_sheets(path: str) -> list[str]:
    """Noms des feuilles d'un classeur (Excel/ODS), ou [] pour les autres formats."""
    ext = Path(path).suffix.lower()
    if ext not in (".xlsx", ".xls", ".ods"):
        return []
    try:
        xls = pd.ExcelFile(path, engine=_excel_engine(ext))
        return [str(s) for s in xls.sheet_names]
    except Exception:  # pragma: no cover - fichier illisible
        return []


def _detect_header_row(raw: pd.DataFrame, limite: int = 30) -> int:
    """Détecte la ligne d'en-tête d'un tableau lu SANS en-tête.

    Porté de la v4.9 desktop : parmi les `limite` premières lignes, on retient la
    PREMIÈRE dont le remplissage (cellules non vides) atteint au moins
    max(3, 0.6 × remplissage_max). Élimine les préambules Kobo/CSPro/Excel.
    """
    n = min(limite, len(raw))
    if n == 0:
        return 0
    remplissages = [int(raw.iloc[i].notna().sum()) for i in range(n)]
    max_fill = max(remplissages) if remplissages else 0
    seuil = max(3, int(0.6 * max_fill))
    for i, f in enumerate(remplissages):
        if f >= seuil:
            return i
    return 0


def _apply_header(raw: pd.DataFrame, header_row: int) -> pd.DataFrame:
    """Applique la ligne `header_row` comme en-tête : noms vides → Colonne_k,
    doublons suffixés, puis retrait des lignes entièrement vides."""
    header_row = max(0, min(header_row, len(raw) - 1)) if len(raw) else 0
    entetes = raw.iloc[header_row].tolist() if len(raw) else []
    noms: list[str] = []
    vus: dict[str, int] = {}
    for k, brut in enumerate(entetes):
        nom = "" if brut is None else str(brut).strip()
        if not nom or nom.lower() == "nan":
            nom = f"Colonne_{k + 1}"
        if nom in vus:
            vus[nom] += 1
            nom = f"{nom}_{vus[nom]}"
        else:
            vus[nom] = 1
        noms.append(nom)
    corps = raw.iloc[header_row + 1 :].copy()
    corps.columns = noms
    corps = corps.dropna(how="all").reset_index(drop=True)
    return corps


def _sniff_csv_sep(path: str, limite: int = 40) -> str:
    """Devine le séparateur d'un CSV en le cherchant sur PLUSIEURS lignes.

    csv.Sniffer se fie à la 1re ligne, ce qui échoue quand un préambule
    Kobo/CSPro (sans séparateur) la précède. On retient, parmi , ; \\t |, celui
    dont le nombre d'occurrences par ligne est le plus élevé et le plus stable
    sur les premières lignes non vides. Par défaut la virgule."""
    candidats = [",", ";", "\t", "|"]
    from collections import Counter

    comptes: dict[str, list[int]] = {c: [] for c in candidats}
    try:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
            for ligne in fh:
                if len(comptes[candidats[0]]) >= limite:
                    break
                if not ligne.strip():
                    continue
                for c in candidats:
                    comptes[c].append(ligne.count(c))
    except OSError:  # pragma: no cover - fichier illisible
        return ","
    meilleur, score = ",", -1.0
    for c in candidats:
        vals = [n for n in comptes[c] if n > 0]
        if not vals:
            continue
        # Fréquence dominante × présence : privilégie un séparateur régulier.
        commun, freq = Counter(vals).most_common(1)[0]
        note = commun * (freq / len(comptes[c])) if comptes[c] else 0
        if note > score:
            meilleur, score = c, note
    return meilleur


def _read_csv_raw(path: str, sep: str) -> pd.DataFrame:
    """Lit un CSV/TSV en DataFrame SANS en-tête, tolérant aux lignes ragged.

    Un préambule Kobo/CSPro a souvent moins de colonnes que les données ; on lit
    ligne par ligne (module csv) et on complète à la largeur maximale, ce qui
    évite l'erreur « Expected N fields » de pandas et préserve la détection
    d'en-tête sur des lignes de tailles différentes."""
    import csv as _csv

    lignes: list[list[Optional[str]]] = []
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
        for cells in _csv.reader(fh, delimiter=sep):
            # Cellule vide → None : notna()/dropna() traitent correctement les
            # préambules et lignes blanches lors de la détection d'en-tête.
            lignes.append([(c if c.strip() != "" else None) for c in cells])
    if not lignes:
        return pd.DataFrame()
    largeur = max(len(row) for row in lignes)
    normalisees = [row + [None] * (largeur - len(row)) for row in lignes]
    return pd.DataFrame(normalisees, dtype=object)


def read_tabular(
    path: str,
    sheet: Optional[str] = None,
    header_row: Optional[int] = None,
) -> SurveyDataset:
    """Lit un fichier tabulaire non-SPSS (Excel Kobo/CSPro, CSV, ODS, Word...).

    - `sheet` : nom de la feuille à analyser (classeurs multi-feuilles) ; 1re par défaut.
    - `header_row` : n° (0-indexé) de la ligne d'en-tête ; auto-détecté si None.

    Aucune métadonnée SPSS n'accompagne ces formats ; les libellés de valeurs
    sont donc vides et les modalités s'affichent avec leur code.
    """
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xls", ".ods"):
        engine = _excel_engine(ext)
        sheets = pd.read_excel(path, sheet_name=None, engine=engine, header=None)
        nom = sheet if (sheet and sheet in sheets) else next(iter(sheets))
        raw = sheets[nom]
        hr = header_row if header_row is not None else _detect_header_row(raw)
        return SurveyDataset(frame=_apply_header(raw, hr), name=str(nom))
    if ext in (".csv", ".tsv", ".tab"):
        sep = "\t" if ext in (".tsv", ".tab") else _sniff_csv_sep(path)
        raw = _read_csv_raw(path, sep)
        hr = header_row if header_row is not None else _detect_header_row(raw)
        return SurveyDataset(frame=_apply_header(raw, hr))
    if ext == ".json":
        return SurveyDataset(frame=pd.read_json(path))
    if ext == ".docx":
        return SurveyDataset(frame=_read_docx_table(path), name=Path(path).stem)
    if ext == ".pbix":
        raise ValueError(
            "Les fichiers Power BI (.pbix) ne sont pas lisibles directement (format "
            "propriétaire compressé). Dans Power BI, exportez les données en CSV ou "
            "Excel (ou via Power Query), puis importez ce fichier."
        )
    if ext == ".doc":
        raise ValueError(
            "L'ancien format Word (.doc) n'est pas lisible. Enregistrez le document "
            "en .docx (avec un tableau) ou exportez les données en Excel/CSV."
        )
    raise ValueError(
        f"Format non pris en charge : {ext}. Formats acceptés : .sav, "
        + ", ".join(TABULAR_EXTENSIONS)
        + "."
    )


def load_dataset(
    path: str,
    sheet: Optional[str] = None,
    header_row: Optional[int] = None,
) -> SurveyDataset:
    """Charge n'importe quel format supporté vers un SurveyDataset.

    `sheet` / `header_row` ne s'appliquent qu'aux formats tabulaires (ignorés
    pour .sav).
    """
    if Path(path).suffix.lower() == ".sav":
        return read_sav(path)
    return read_tabular(path, sheet=sheet, header_row=header_row)


def write_sav(dataset: SurveyDataset, frame: pd.DataFrame, path: str, specs: Optional[dict] = None) -> None:
    """Écrit `frame` en .sav en conservant étiquettes et niveaux de mesure SPSS."""
    if specs is None:
        specs = infer_variable_specs(dataset)
    measures = {
        c: ("scale" if sp["measure"] == "ÉCHELLE" else sp["measure"].lower())
        for c, sp in specs.items()
        if c in frame.columns
    }
    formats = {
        c: ("F12.0" if sp["decimals"] == 0 else "F12.2")
        for c, sp in specs.items()
        if c in frame.columns and sp["decimals"] is not None
    }
    _pyreadstat().write_sav(
        frame,
        path,
        column_labels=dataset.variable_labels or None,
        variable_value_labels=dataset.value_labels or None,
        variable_measure=measures or None,
        variable_format=formats or None,
    )
